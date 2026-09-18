from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import PermissionRequiredMixin
from django.core.paginator import Paginator
from django.db.models import Count, Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse_lazy
from django.utils import timezone
from django.views.generic import CreateView, DeleteView, ListView, UpdateView

from .forms import ContactForm
from .models import Contact


def permission_denied(request, exception=None):
    return render(request, "403.html", status=403)


def get_kpis():
    active = Contact.objects.filter(status=Contact.Status.ACTIVE)
    tier_a = active.filter(tier=Contact.Tier.TIER_A)
    total_active = active.count()
    overdue = active.filter(status=Contact.Status.ACTIVE)
    overdue = [c for c in active.all() if c.is_validation_overdue]

    tier_a_total = tier_a.count()
    tier_a_valid = sum(1 for c in tier_a.all() if not c.is_validation_overdue)
    active_valid = sum(1 for c in active.all() if not c.is_validation_overdue)
    unowned = Contact.objects.filter(owner__isnull=True).count()

    kpis = {
        "tier_a_total": tier_a_total,
        "tier_a_valid": tier_a_valid,
        "tier_a_pct": round((tier_a_valid / tier_a_total * 100) if tier_a_total else 0, 1),
        "overall_pct": round((active_valid / total_active * 100) if total_active else 0, 1),
        "total_active": total_active,
        "overdue_count": len(overdue),
        "overdue_pct": round((len(overdue) / total_active * 100) if total_active else 0, 1),
        "overdue_tier_a": sum(1 for c in overdue if c.tier == Contact.Tier.TIER_A),
        "unowned": unowned,
        "tier_counts": {t: active.filter(tier=t).count() for t in Contact.Tier.values},
        "tier_validation": {
            t: {
                "total": active.filter(tier=t).count(),
                "validated": sum(
                    1
                    for c in active.filter(tier=t).all()
                    if not c.is_validation_overdue
                ),
                "overdue": sum(
                    1 for c in active.filter(tier=t).all() if c.is_validation_overdue
                ),
            }
            for t in Contact.Tier.values
        },
        "status_counts": {
            s: Contact.objects.filter(status=s).count() for s in Contact.Status.values
        },
        "unresolved_duplicates": len(find_duplicates()),
    }
    return kpis


def find_duplicates():
    """Return list of dicts for records sharing the same name+company key."""
    groups = {}
    for c in Contact.objects.all():
        key = (c.first_name.strip().lower(), c.last_name.strip().lower(), c.company.strip().lower())
        groups.setdefault(key, []).append(c)
    return [
        {
            "key": f"{g[0].first_name} {g[0].last_name} — {g[0].company}",
            "names": [c.full_name for c in g],
            "ids": [c.id for c in g],
        }
        for g in groups.values()
        if len(g) > 1
    ]


def owner_breakdown():
    """Contact counts, completion %, overdue per Account Manager."""
    rows = []
    for item in (
        Contact.objects.exclude(owner__isnull=True)
        .filter(status=Contact.Status.ACTIVE)
        .values("owner_id", "owner__first_name", "owner__last_name", "owner__email")
        .annotate(total=Count("id"))
        .order_by("owner__first_name")
    ):
        ids = [
            c.id
            for c in Contact.objects.filter(
                owner_id=item["owner_id"], status=Contact.Status.ACTIVE
            ).all()
            if c.is_validation_overdue
        ]
        rows.append(
            {
                "name": f"{item['owner__first_name']} {item['owner__last_name']}",
                "email": item["owner__email"],
                "total": item["total"],
                "overdue": len(ids),
                "completion": round((item["total"] - len(ids)) / item["total"] * 100)
                if item["total"]
                else 0,
            }
        )
    rows.sort(key=lambda r: r["overdue"], reverse=True)
    return rows


@login_required
def dashboard(request):
    kpis = get_kpis()
    tiers = []
    for label in ["A", "B", "C"]:
        tv = kpis["tier_validation"][label]
        tiers.append(
            {
                "label": label,
                "total": tv["total"],
                "validated": tv["validated"],
                "overdue": tv["overdue"],
                "pct": round((tv["validated"] / tv["total"] * 100) if tv["total"] else 0, 1),
            }
        )
    from django.utils import timezone

    hour = timezone.localtime(timezone.now()).hour
    if hour < 12:
        greeting = "Good morning"
    elif hour < 18:
        greeting = "Good afternoon"
    else:
        greeting = "Good evening"
    return render(
        request,
        "contacts/dashboard.html",
        {
            "kpis": kpis,
            "tiers": tiers,
            "owners": owner_breakdown(),
            "greeting": greeting,
            "today": timezone.localtime(timezone.now()),
        },
    )


class ContactListView(PermissionRequiredMixin, ListView):
    model = Contact
    template_name = "contacts/contact_list.html"
    context_object_name = "contacts"
    paginate_by = 25
    permission_required = "contacts.view_contact"

    def get_queryset(self):
        qs = super().get_queryset().select_related("owner")
        q = self.request.GET.get("q")
        tier = self.request.GET.get("tier")
        status = self.request.GET.get("status")
        if q:
            qs = qs.filter(
                Q(first_name__icontains=q)
                | Q(last_name__icontains=q)
                | Q(company__icontains=q)
                | Q(position__icontains=q)
                | Q(email__icontains=q)
                | Q(owner__username__icontains=q)
                | Q(owner__first_name__icontains=q)
                | Q(owner__last_name__icontains=q)
            )
        if tier:
            qs = qs.filter(tier=tier)
        if status:
            qs = qs.filter(status=status)
        return qs

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx["tiers"] = Contact.Tier.choices
        ctx["statuses"] = Contact.Status.choices
        return ctx


class MyContactsView(PermissionRequiredMixin, ListView):
    model = Contact
    template_name = "contacts/contact_list.html"
    context_object_name = "contacts"
    paginate_by = 25
    permission_required = "contacts.view_contact"

    def get_queryset(self):
        qs = Contact.objects.filter(owner=self.request.user).select_related("owner")
        if self.request.GET.get("overdue") == "1":
            qs = qs.filter(status=Contact.Status.ACTIVE)
            ids = [c.id for c in qs.all() if c.is_validation_overdue]
            qs = Contact.objects.filter(id__in=ids)
        return qs

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx["is_my_contacts"] = True
        return ctx


class ContactCreateView(PermissionRequiredMixin, CreateView):
    model = Contact
    form_class = ContactForm
    template_name = "contacts/contact_form.html"
    success_url = reverse_lazy("contact_list")
    permission_required = "contacts.add_contact"

    def get_initial(self):
        return {"owner": self.request.user}


class ContactUpdateView(PermissionRequiredMixin, UpdateView):
    model = Contact
    form_class = ContactForm
    template_name = "contacts/contact_form.html"
    success_url = reverse_lazy("contact_list")
    permission_required = "contacts.change_contact"


class ContactDeleteView(PermissionRequiredMixin, DeleteView):
    model = Contact
    success_url = reverse_lazy("contact_list")
    permission_required = "contacts.delete_contact"


@login_required
def validate_contact(request, pk):
    contact = get_object_or_404(Contact, pk=pk)
    if contact.owner != request.user and not request.user.is_superuser:
        from django.contrib.auth.models import Group

        if not request.user.groups.filter(name="Commercial Manager").exists():
            messages.error(request, "Only the owner or custodian may validate this contact.")
            return redirect("contact_list")
    contact.last_validated_at = timezone.now()
    contact.status = Contact.Status.ACTIVE
    contact.save()
    messages.success(request, f"{contact.full_name} validated.")
    return redirect(request.META.get("HTTP_REFERER", "contact_list"))


@login_required
def events(request):
    contacts = Contact.objects.filter(status=Contact.Status.ACTIVE).select_related("owner")
    q = request.GET.get("q")
    tier = request.GET.get("tier")
    company = request.GET.get("company")
    if q:
        contacts = contacts.filter(
            Q(first_name__icontains=q)
            | Q(last_name__icontains=q)
            | Q(company__icontains=q)
            | Q(email__icontains=q)
            | Q(owner__username__icontains=q)
            | Q(owner__first_name__icontains=q)
            | Q(owner__last_name__icontains=q)
        )
    if tier:
        contacts = contacts.filter(tier=tier)
    if company:
        contacts = contacts.filter(company__icontains=company)
    return render(
        request,
        "contacts/events.html",
        {
            "contacts": contacts,
            "tiers": Contact.Tier.choices,
            "picked_tier": tier,
            "picked_company": company,
            "q": q,
        },
    )


@login_required
def export_csv(request):
    contacts = Contact.objects.filter(status=Contact.Status.ACTIVE).select_related("owner")
    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = 'attachment; filename="guest_list.csv"'
    import csv

    writer = csv.writer(response)
    writer.writerow(
        [
            "Full Name",
            "Position",
            "Company",
            "Email",
            "Telephone",
            "Relationship Owner",
            "Strategic Tier",
            "Last Validation Date",
            "Status",
        ]
    )
    for c in contacts:
        writer.writerow(
            [
                c.full_name,
                c.position,
                c.company,
                c.email,
                c.telephone,
                c.owner.get_full_name() or c.owner.username,
                c.tier,
                c.last_validated_at.date().isoformat(),
                c.status,
            ]
        )
    return response


@login_required
def export_xlsx(request):
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    contacts = Contact.objects.filter(status=Contact.Status.ACTIVE).select_related("owner")
    wb = Workbook()
    ws = wb.active
    ws.title = "Contacts"

    headers = [
        "Name", "Position", "Company", "Email", "Telephone",
        "Relationship Owner", "Strategic Tier", "Last Validation Date",
        "Status", "Validation Overdue",
    ]
    ws.append(headers)
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font

    for c in contacts:
        ws.append(
            [
                c.full_name, c.position, c.company, c.email, c.telephone,
                c.owner.get_full_name() or c.owner.username, f"Tier {c.tier}",
                c.last_validated_at.date().isoformat(), c.get_status_display(),
                "Yes" if c.is_validation_overdue else "No",
            ]
        )

    for i, w in enumerate([28, 30, 30, 32, 20, 30, 12, 20, 12, 18], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="guest_database.xlsx"'
    return response


@login_required
def duplicate_check(request):
    return render(request, "contacts/duplicates.html", {"duplicates": find_duplicates()})